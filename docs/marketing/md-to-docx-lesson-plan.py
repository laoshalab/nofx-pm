#!/usr/bin/env python3
"""Convert prediction-video-lesson-plan.md to .docx (minimal deps: python-docx)."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn


def set_cn_font(run, name="Microsoft YaHei", size_pt=11):
    run.font.name = name
    run.font.size = Pt(size_pt)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def add_rich_paragraph(doc, text, style=None, base_size=11, bold_all=False):
    p = doc.add_paragraph(style=style)
    if bold_all:
        run = p.add_run(text)
        run.bold = True
        set_cn_font(run, size_pt=base_size)
        return p

    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
            set_cn_font(run, size_pt=base_size)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        else:
            run = p.add_run(part)
            set_cn_font(run, size_pt=base_size)
    return p


def parse_table_lines(lines):
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            break
        if re.match(r"^\|[-:\s|]+\|$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
    return rows


def md_to_docx(md_path: Path, docx_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("## "):
            t = doc.add_heading(stripped[2:].strip(), level=0)
            for run in t.runs:
                set_cn_font(run, size_pt=18)
            i += 1
            continue

        if stripped.startswith("## "):
            t = doc.add_heading(stripped[3:].strip(), level=1)
            for run in t.runs:
                set_cn_font(run, size_pt=14)
            i += 1
            continue

        if stripped.startswith("### "):
            t = doc.add_heading(stripped[4:].strip(), level=2)
            for run in t.runs:
                set_cn_font(run, size_pt=12)
            i += 1
            continue

        if stripped.startswith("> "):
            block = []
            while i < len(lines) and lines[i].strip().startswith("> "):
                block.append(lines[i].strip()[2:].strip())
                i += 1
            add_rich_paragraph(doc, " ".join(block), base_size=10)
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = parse_table_lines(table_lines)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        cell_para = table.rows[ri].cells[ci].paragraphs[0]
                        cell_para.clear()
                        run = cell_para.add_run(cell)
                        set_cn_font(run, size_pt=10)
                        if ri == 0:
                            run.bold = True
            continue

        if stripped.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                add_rich_paragraph(doc, lines[i].strip()[2:], style="List Bullet", base_size=10)
                i += 1
            continue

        if re.match(r"^\*\*[^*]+：\*\*", stripped) or stripped.startswith("**第"):
            add_rich_paragraph(doc, stripped, base_size=11)
            i += 1
            continue

        add_rich_paragraph(doc, stripped, base_size=11)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


if __name__ == "__main__":
    base = Path(__file__).resolve().parent
    md = base / "prediction-video-lesson-plan.md"
    out = base / "prediction-video-lesson-plan.docx"
    if len(sys.argv) > 1:
        md = Path(sys.argv[1])
    if len(sys.argv) > 2:
        out = Path(sys.argv[2])
    md_to_docx(md, out)
    print(out)
